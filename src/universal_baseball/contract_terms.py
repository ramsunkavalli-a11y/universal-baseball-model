"""FanGraphs payroll normalization for terms StatsAPI does not provide."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from typing import Mapping

import polars as pl


PLAYER_SECTIONS = frozenset(
    {
        "Guaranteed",
        "Eligible For Arb",
        "Not Yet Eligible For Arb",
        "No Longer On 40-Man Roster",
    }
)

PAYROLL_PLAYER_SCHEMA: dict[str, pl.DataType] = {
    "season": pl.Int64,
    "team_name": pl.String,
    "payroll_section": pl.String,
    "fangraphs_id": pl.String,
    "player_name": pl.String,
    "reference_service_time": pl.String,
    "contract_text": pl.String,
    "aav_dollars": pl.Int64,
    "clause_parse_status": pl.String,
    "source_snapshot_id": pl.String,
}

PAYROLL_YEAR_SCHEMA: dict[str, pl.DataType] = {
    "season": pl.Int64,
    "team_name": pl.String,
    "payroll_section": pl.String,
    "fangraphs_id": pl.String,
    "player_name": pl.String,
    "payroll_year": pl.Int64,
    "raw_value": pl.String,
    "amount_dollars": pl.Int64,
    "term_label": pl.String,
    "source_snapshot_id": pl.String,
}

CONTRACT_CLAUSE_SCHEMA: dict[str, pl.DataType] = {
    "fangraphs_id": pl.String,
    "player_name": pl.String,
    "clause_type": pl.String,
    "start_year": pl.Int64,
    "end_year": pl.Int64,
    "raw_contract_text": pl.String,
    "source_snapshot_id": pl.String,
}

OTHER_PAYMENT_SCHEMA: dict[str, pl.DataType] = {
    "season": pl.Int64,
    "team_name": pl.String,
    "payment_year": pl.Int64,
    "description": pl.String,
    "amount_dollars": pl.Int64,
    "payment_type": pl.String,
    "is_contingent": pl.Boolean,
    "source_snapshot_id": pl.String,
}


@dataclass(frozen=True)
class PayrollNormalization:
    players: pl.DataFrame
    year_terms: pl.DataFrame
    clauses: pl.DataFrame
    other_payments: pl.DataFrame


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _money(value: object) -> int | None:
    text = _text(value)
    if not text.startswith(("$", "-$")):
        return None
    negative = text.startswith("-")
    digits = text.replace("-$", "").replace("$", "").replace(",", "")
    try:
        amount = int(round(float(digits)))
    except ValueError as exc:
        raise ValueError(f"invalid payroll money value: {text}") from exc
    return -amount if negative else amount


def _service_time(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return f"{float(value):.3f}"
    return _text(value)


def _term_label(raw: str, amount: int | None) -> str:
    if amount is not None:
        return "salary_or_option_amount"
    normalized = raw.upper()
    if normalized == "FREE AGENT":
        return "free_agent"
    if normalized == "PRE-ARB":
        return "pre_arbitration"
    if normalized.startswith("ARB ") and normalized[4:].isdigit():
        return f"arbitration_{normalized[4:]}"
    if normalized == "TBD":
        return "to_be_determined"
    return "unparsed"


OPTION_PATTERN = re.compile(
    r"(?P<start>20\d{2})(?:-(?P<end>\d{2}|20\d{2}))?\s+"
    r"(?P<kind>club|player|mutual)\s+options?",
    re.IGNORECASE,
)
OPT_OUT_YEAR_PATTERN = re.compile(r"(?:can|may) opt out after (?P<year>20\d{2})", re.IGNORECASE)


def _full_year(start: int, end_text: str | None) -> int:
    if end_text is None:
        return start
    end = int(end_text)
    return end if end >= 2000 else (start // 100) * 100 + end


def _clauses(
    *, player_name: str, fangraphs_id: str, contract_text: str, source_snapshot_id: str
) -> tuple[list[dict[str, object]], str]:
    rows: list[dict[str, object]] = []
    lowered = contract_text.lower()
    for match in OPTION_PATTERN.finditer(contract_text):
        start = int(match.group("start"))
        end = _full_year(start, match.group("end"))
        kind = match.group("kind").lower()
        clause_type = f"{kind}_option"
        if kind == "club" and f"{start} club option exercised" in lowered:
            clause_type = "exercised_club_option"
        rows.append(
            {
                "fangraphs_id": fangraphs_id,
                "player_name": player_name,
                "clause_type": clause_type,
                "start_year": start,
                "end_year": end,
                "raw_contract_text": contract_text,
                "source_snapshot_id": source_snapshot_id,
            }
        )
    for match in OPT_OUT_YEAR_PATTERN.finditer(contract_text):
        year = int(match.group("year"))
        rows.append(
            {
                "fangraphs_id": fangraphs_id,
                "player_name": player_name,
                "clause_type": "opt_out_after_year",
                "start_year": year,
                "end_year": year,
                "raw_contract_text": contract_text,
                "source_snapshot_id": source_snapshot_id,
            }
        )
    if "opt out after each year" in lowered:
        rows.append(
            {
                "fangraphs_id": fangraphs_id,
                "player_name": player_name,
                "clause_type": "opt_out_after_each_year",
                "start_year": None,
                "end_year": None,
                "raw_contract_text": contract_text,
                "source_snapshot_id": source_snapshot_id,
            }
        )
    expected = any(term in lowered for term in ("option", "opt out"))
    return rows, "review_unparsed_clause" if expected and not rows else "parsed_or_not_applicable"


def _year_columns(frame: pl.DataFrame) -> list[str]:
    return sorted(
        [column for column in frame.columns if str(column).isdigit()],
        key=int,
    )


def normalize_fangraphs_payroll(
    sheets: Mapping[str, pl.DataFrame],
    *,
    team_name: str,
    season: int,
    source_snapshot_id: str,
) -> PayrollNormalization:
    """Normalize player contracts, annual terms, clauses and other payments."""

    if not team_name.strip() or not source_snapshot_id.strip():
        raise ValueError("team_name and source_snapshot_id must be nonblank")
    missing_sections = sorted(PLAYER_SECTIONS - set(sheets))
    if missing_sections:
        raise ValueError(f"payroll workbook missing player sections: {missing_sections}")
    player_rows: list[dict[str, object]] = []
    year_rows: list[dict[str, object]] = []
    clause_rows: list[dict[str, object]] = []
    for section in sorted(PLAYER_SECTIONS):
        frame = sheets[section].rename({column: str(column) for column in sheets[section].columns})
        required = {"Player", "Service Time", "Contract", "AAV", "playerId"}
        missing = sorted(required - set(frame.columns))
        if missing:
            raise ValueError(f"payroll section {section} missing columns: {missing}")
        for row in frame.iter_rows(named=True):
            player_name = _text(row["Player"])
            fangraphs_id = _text(row["playerId"])
            if not player_name or not fangraphs_id:
                raise ValueError(f"payroll section {section} has missing player identity")
            contract_text = _text(row["Contract"])
            parsed_clauses, clause_status = _clauses(
                player_name=player_name,
                fangraphs_id=fangraphs_id,
                contract_text=contract_text,
                source_snapshot_id=source_snapshot_id,
            )
            clause_rows.extend(parsed_clauses)
            player_rows.append(
                {
                    "season": season,
                    "team_name": team_name,
                    "payroll_section": section,
                    "fangraphs_id": fangraphs_id,
                    "player_name": player_name,
                    "reference_service_time": _service_time(row["Service Time"]),
                    "contract_text": contract_text,
                    "aav_dollars": _money(row["AAV"]),
                    "clause_parse_status": clause_status,
                    "source_snapshot_id": source_snapshot_id,
                }
            )
            for year_column in _year_columns(frame):
                raw = _text(row[year_column])
                if not raw:
                    continue
                amount = _money(raw)
                year_rows.append(
                    {
                        "season": season,
                        "team_name": team_name,
                        "payroll_section": section,
                        "fangraphs_id": fangraphs_id,
                        "player_name": player_name,
                        "payroll_year": int(year_column),
                        "raw_value": raw,
                        "amount_dollars": amount,
                        "term_label": _term_label(raw, amount),
                        "source_snapshot_id": source_snapshot_id,
                    }
                )

    payment_rows: list[dict[str, object]] = []
    if "Other Payments" in sheets:
        payments = sheets["Other Payments"].rename(
            {column: str(column) for column in sheets["Other Payments"].columns}
        )
        if "Description" not in payments.columns:
            raise ValueError("Other Payments missing Description")
        for row in payments.iter_rows(named=True):
            description = _text(row["Description"])
            lowered = description.lower()
            for year_column in _year_columns(payments):
                amount = _money(row[year_column])
                if amount is None:
                    continue
                payment_type = (
                    "buyout"
                    if "buyout" in lowered
                    else "trade_credit"
                    if amount < 0 and any(word in lowered for word in ("paid by", "owed from"))
                    else "other"
                )
                payment_rows.append(
                    {
                        "season": season,
                        "team_name": team_name,
                        "payment_year": int(year_column),
                        "description": description,
                        "amount_dollars": amount,
                        "payment_type": payment_type,
                        "is_contingent": "potential" in lowered,
                        "source_snapshot_id": source_snapshot_id,
                    }
                )

    return PayrollNormalization(
        players=pl.DataFrame(player_rows, schema=PAYROLL_PLAYER_SCHEMA).sort("fangraphs_id"),
        year_terms=pl.DataFrame(year_rows, schema=PAYROLL_YEAR_SCHEMA).sort(
            ["fangraphs_id", "payroll_year"]
        ),
        clauses=(
            pl.DataFrame(clause_rows, schema=CONTRACT_CLAUSE_SCHEMA)
            if clause_rows
            else pl.DataFrame(schema=CONTRACT_CLAUSE_SCHEMA)
        ).sort(["fangraphs_id", "start_year"]),
        other_payments=(
            pl.DataFrame(payment_rows, schema=OTHER_PAYMENT_SCHEMA)
            if payment_rows
            else pl.DataFrame(schema=OTHER_PAYMENT_SCHEMA)
        ).sort(["payment_year", "description"]),
    )


def read_fangraphs_payroll_xlsx(path: Path) -> dict[str, pl.DataFrame]:
    """Read a payroll workbook; normalization remains a separate audited step."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install the project payroll extra to read xlsx files") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, pl.DataFrame] = {}
    for worksheet in workbook.worksheets:
        values = list(worksheet.values)
        if not values:
            continue
        headers = [str(value) if value is not None else f"unnamed_{index}" for index, value in enumerate(values[0])]
        rows = [dict(zip(headers, row, strict=True)) for row in values[1:] if any(value is not None for value in row)]
        sheets[worksheet.title] = pl.DataFrame(rows, infer_schema_length=None)
    workbook.close()
    return sheets
