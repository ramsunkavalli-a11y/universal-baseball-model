"""FanGraphs Opening Day Tracker adapter for historical control baselines."""

from __future__ import annotations

from dataclasses import dataclass
import json
from pathlib import Path
import re
from typing import Any, Iterable, Mapping

import polars as pl

from universal_baseball.control_baseline import service_time_to_days


OPENING_DAY_CONTROL_SCHEMA: dict[str, pl.DataType] = {
    "season": pl.Int64,
    "player_id": pl.Int64,
    "fangraphs_id": pl.String,
    "player_name": pl.String,
    "team_abbreviation": pl.String,
    "fangraphs_team_id": pl.Int64,
    "service_time": pl.String,
    "service_days": pl.Int64,
    "options_remaining": pl.Int64,
    "rule5_status": pl.String,
    "rule5_eligibility_year": pl.Int64,
    "on_40man": pl.Boolean,
    "roster_status": pl.String,
    "projected_opening_day_role": pl.String,
    "source_snapshot_id": pl.String,
}

OPENING_DAY_PROJECTION_SCHEMA: dict[str, pl.DataType] = {
    "season": pl.Int64,
    "player_id": pl.Int64,
    "fangraphs_id": pl.String,
    "player_name": pl.String,
    "team_abbreviation": pl.String,
    "position": pl.String,
    "age": pl.Float64,
    "projected_pa": pl.Float64,
    "projected_ip": pl.Float64,
    "projected_opening_day_role": pl.String,
    "source_snapshot_id": pl.String,
}

OPENING_DAY_WORKBOOK_COLUMNS = frozenset(
    {
        "Team",
        "Name",
        "Pos",
        "Age",
        "Proj PA",
        "Proj IP",
        "Service Time",
        "Options or R5 Status",
        "Projected Opening Day Status",
        "PlayerId",
        "MLBAMID",
    }
)


@dataclass(frozen=True, slots=True)
class OpeningDayWorkbookNormalization:
    control: pl.DataFrame
    projections: pl.DataFrame

NEXT_DATA_PREFIX = '<script id="__NEXT_DATA__" type="application/json">'


def extract_next_data_document(page_html: str) -> dict[str, Any]:
    """Extract the server-rendered Next.js data document without browser state."""

    start = page_html.find(NEXT_DATA_PREFIX)
    if start < 0:
        raise ValueError("FanGraphs page lacks __NEXT_DATA__")
    start += len(NEXT_DATA_PREFIX)
    end = page_html.find("</script>", start)
    if end < 0:
        raise ValueError("FanGraphs page has unterminated __NEXT_DATA__")
    payload = json.loads(page_html[start:end])
    if not isinstance(payload, dict):
        raise ValueError("FanGraphs __NEXT_DATA__ must be an object")
    return payload


def _tracker_rows(document: dict[str, Any], season: int) -> list[dict[str, Any]]:
    try:
        queries = document["props"]["pageProps"]["dehydratedState"]["queries"]
    except (KeyError, TypeError) as exc:
        raise ValueError("FanGraphs page lacks dehydrated tracker data") from exc
    matches = []
    for query in queries:
        key = query.get("queryKey") if isinstance(query, dict) else None
        if (
            isinstance(key, list)
            and len(key) == 2
            and key[0] == "roster-resource/opening-day-tracker/data"
            and isinstance(key[1], dict)
            and int(key[1].get("season", -1)) == season
        ):
            matches.append(query)
    if len(matches) != 1:
        raise ValueError("FanGraphs page must contain one matching tracker query")
    rows = matches[0].get("state", {}).get("data")
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise ValueError("FanGraphs tracker data must be a list of objects")
    return rows


def _control_reference(value: object) -> tuple[int | None, str, int | None]:
    if value in (None, "", "n/a", "N/A"):
        return None, "", None
    text = str(value).strip()
    if text.isdigit() and int(text) <= 5:
        return int(text), "", None
    if text.upper() == "R5":
        return None, "rule5_eligible", None
    match = re.fullmatch(r"Dec'(\d{2})", text, re.IGNORECASE)
    if match:
        return None, "rule5_not_yet_eligible", 2000 + int(match.group(1))
    raise ValueError(f"invalid FanGraphs option/Rule 5 value: {value}")


def project_opening_day_control_baseline(
    document: dict[str, Any], *, season: int, source_snapshot_id: str
) -> pl.DataFrame:
    """Return one MLBAM-keyed pre-season service/control reference per player."""

    projected: list[dict[str, object]] = []
    for row in _tracker_rows(document, season):
        if int(row.get("season", -1)) != season:
            raise ValueError("FanGraphs tracker row has wrong season")
        if row.get("xMLBAMID") is None:
            raise ValueError("FanGraphs tracker row lacks MLBAM identity")
        service_time = str(row.get("servicetime") or "").strip()
        options, rule5_status, rule5_year = _control_reference(row.get("options"))
        projected.append(
            {
                "season": season,
                "player_id": int(row["xMLBAMID"]),
                "fangraphs_id": str(row.get("playerId") or "").strip(),
                "player_name": str(row.get("playerName") or "").strip(),
                "team_abbreviation": str(row.get("team") or "").strip(),
                "fangraphs_team_id": int(row["playerTeamId"])
                if row.get("playerTeamId") is not None
                else None,
                "service_time": service_time,
                "service_days": service_time_to_days(service_time),
                "options_remaining": options,
                "rule5_status": rule5_status,
                "rule5_eligibility_year": rule5_year,
                "on_40man": str(row.get("is40Man") or "").upper() == "Y",
                "roster_status": str(row.get("status") or "").strip(),
                "projected_opening_day_role": str(
                    row.get("projectedOpeningDayRole") or ""
                ).strip(),
                "source_snapshot_id": source_snapshot_id,
            }
        )

    frame = pl.DataFrame(projected, schema=OPENING_DAY_CONTROL_SCHEMA)
    collapsed: list[dict[str, object]] = []
    compare = [column for column in OPENING_DAY_CONTROL_SCHEMA if column != "projected_opening_day_role"]
    for group in frame.partition_by("player_id", maintain_order=True):
        if group.select(compare).unique().height != 1:
            raise ValueError("FanGraphs tracker has conflicting duplicate MLBAM rows")
        row = group.row(0, named=True)
        row["projected_opening_day_role"] = ",".join(
            sorted(set(group.get_column("projected_opening_day_role").to_list()))
        )
        collapsed.append(row)
    return pl.DataFrame(collapsed, schema=OPENING_DAY_CONTROL_SCHEMA).sort("player_id")


def _optional_float(value: object) -> float | None:
    if value in (None, ""):
        return None
    return float(value)


def project_opening_day_workbook_rows(
    rows: Iterable[Mapping[str, object]],
    *,
    season: int,
    source_snapshot_id: str,
) -> OpeningDayWorkbookNormalization:
    """Normalize a member workbook without inferring fields it does not contain."""

    control_rows: list[dict[str, object]] = []
    projection_rows: list[dict[str, object]] = []
    for row in rows:
        missing = sorted(OPENING_DAY_WORKBOOK_COLUMNS - set(row))
        if missing:
            raise ValueError(f"FanGraphs Opening Day workbook missing columns: {missing}")
        if row["MLBAMID"] in (None, ""):
            raise ValueError("FanGraphs Opening Day workbook row lacks MLBAM identity")
        player_id = int(row["MLBAMID"])
        service_time = str(row["Service Time"] or "").strip()
        options, rule5_status, rule5_year = _control_reference(
            row["Options or R5 Status"]
        )
        common = {
            "season": season,
            "player_id": player_id,
            "fangraphs_id": str(row["PlayerId"] or "").strip(),
            "player_name": str(row["Name"] or "").strip(),
            "team_abbreviation": str(row["Team"] or "").strip(),
            "projected_opening_day_role": str(
                row["Projected Opening Day Status"] or ""
            ).strip(),
            "source_snapshot_id": source_snapshot_id,
        }
        control_rows.append(
            {
                **common,
                "fangraphs_team_id": None,
                "service_time": service_time,
                "service_days": service_time_to_days(service_time),
                "options_remaining": options,
                "rule5_status": rule5_status,
                "rule5_eligibility_year": rule5_year,
                "on_40man": None,
                "roster_status": "",
            }
        )
        projection_rows.append(
            {
                **common,
                "position": str(row["Pos"] or "").strip(),
                "age": _optional_float(row["Age"]),
                "projected_pa": _optional_float(row["Proj PA"]),
                "projected_ip": _optional_float(row["Proj IP"]),
            }
        )

    control = pl.DataFrame(control_rows, schema=OPENING_DAY_CONTROL_SCHEMA)
    projections = pl.DataFrame(projection_rows, schema=OPENING_DAY_PROJECTION_SCHEMA)
    control_compare = list(OPENING_DAY_CONTROL_SCHEMA)
    projection_compare = list(OPENING_DAY_PROJECTION_SCHEMA)
    collapsed_control: list[dict[str, object]] = []
    collapsed_projections: list[dict[str, object]] = []
    for group in control.partition_by("player_id", maintain_order=True):
        if group.select(control_compare).unique().height != 1:
            raise ValueError("FanGraphs workbook has conflicting duplicate control rows")
        collapsed_control.append(group.row(0, named=True))
    for group in projections.partition_by("player_id", maintain_order=True):
        if group.select(projection_compare).unique().height != 1:
            raise ValueError("FanGraphs workbook has conflicting duplicate projection rows")
        collapsed_projections.append(group.row(0, named=True))
    return OpeningDayWorkbookNormalization(
        control=pl.DataFrame(
            collapsed_control, schema=OPENING_DAY_CONTROL_SCHEMA
        ).sort("player_id"),
        projections=pl.DataFrame(
            collapsed_projections, schema=OPENING_DAY_PROJECTION_SCHEMA
        ).sort("player_id"),
    )


def read_opening_day_tracker_xlsx(
    path: Path, *, season: int, source_snapshot_id: str
) -> OpeningDayWorkbookNormalization:
    """Read one downloaded FanGraphs workbook into audited control/projection tables."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install the project payroll extra to read xlsx files") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Opening Day Tracker" not in workbook.sheetnames:
        workbook.close()
        raise ValueError("FanGraphs workbook lacks Opening Day Tracker sheet")
    worksheet = workbook["Opening Day Tracker"]
    values = worksheet.iter_rows(values_only=True)
    headers = next(values, None)
    if headers is None:
        workbook.close()
        raise ValueError("FanGraphs Opening Day workbook is empty")
    header_names = [str(value) if value is not None else "" for value in headers]
    rows = [
        dict(zip(header_names, values_row, strict=True))
        for values_row in values
        if any(value is not None for value in values_row)
    ]
    workbook.close()
    return project_opening_day_workbook_rows(
        rows, season=season, source_snapshot_id=source_snapshot_id
    )
